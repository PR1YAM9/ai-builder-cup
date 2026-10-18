"""Fill the AI Builder Cup submission template with ClaimGuard content.

Run from deck/:  python build_deck.py   -> ClaimGuard_AI_Builder_Cup.pptx
Keeps every template slide's header, logos and title box; adds content on top.
"""

import copy
import json
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
SHOTS = HERE / "shots"
IMG = HERE / "img"
IMG.mkdir(exist_ok=True)

LIVE = "https://claimguard-ai-342329356911.us-central1.run.app"
REPO = "https://github.com/PR1YAM9/ai-builder-cup"
VIDEO = "Add your video link here"

INK = RGBColor(0x1C, 0x1D, 0x1F)
BODY = RGBColor(0x3D, 0x3F, 0x43)
MUTED = RGBColor(0x6E, 0x70, 0x76)
RULE = RGBColor(0xDD, 0xDF, 0xE3)
CARD = RGBColor(0xF4, 0xF5, 0xF7)
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
RED = RGBColor(0x9A, 0x3B, 0x2E)
GREEN = RGBColor(0x3D, 0x6B, 0x4F)
AMBER = RGBColor(0x8A, 0x6A, 0x1F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"

L, R = 0.35, 9.65          # content margins (inches)
W = R - L


# ---------- helpers ----------
def text(slide, x, y, w, h, paras, size=12, color=BODY, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space_after=4, line=1.15):
    """paras: str, or list of str / (str, dict) / list-of-runs [(str, dict), ...]."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.space_after = Pt(space_after)
        para.line_spacing = line
        runs = p if isinstance(p, list) else [p]
        for r in runs:
            t, o = (r, {}) if isinstance(r, str) else r
            run = para.add_run()
            run.text = t
            f = run.font
            f.name = FONT
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", False)
            f.color.rgb = o.get("color", color)
    return tb


def box(slide, x, y, w, h, fill=CARD, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.text_frame.text = ""
    return s


def boxed(slide, x, y, w, h, paras, fill=CARD, line=None, size=11, color=BODY, pad=0.14,
          anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, radius=0.06):
    box(slide, x, y, w, h, fill=fill, line=line, radius=radius)
    return text(slide, x + pad, y + pad, w - 2 * pad, h - 2 * pad, paras, size=size, color=color,
                anchor=anchor, align=align)


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=1.25):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return c


def picture(slide, path, x, y, w=None, h=None, border=True):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), **kw)
    if border:
        pic.line.color.rgb = RULE
        pic.line.width = Pt(0.75)
    return pic


def crop(src, box_px, name, max_w=2200):
    im = Image.open(src)
    im = im.crop(box_px)
    if im.width > max_w:
        im = im.resize((max_w, int(im.height * max_w / im.width)), Image.LANCZOS)
    out = IMG / name
    im.save(out, optimize=True)
    return out


def set_title(slide, label):
    """Keep the template title box and its formatting; keep only its first paragraph."""
    tb = next(s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip())
    tf = tb.text_frame
    first = tf.paragraphs[0]
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    runs = first.runs
    runs[0].text = label
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    # strip any line-break/vertical-tab elements after the label
    for br in first._p.findall(qn("a:br")):
        first._p.remove(br)
    tb.height = Inches(0.4)
    return tb


def headline(slide, s, y=1.1):
    return text(slide, L, y, W, 0.5, s, size=20, color=INK, bold=True, line=1.05)


def caption(slide, x, y, w, s, align=PP_ALIGN.LEFT):
    return text(slide, x, y, w, 0.3, s, size=10, color=MUTED, align=align)


def delete_slide(prs, slide):
    sld_ids = prs.slides._sldIdLst
    for sld_id in list(sld_ids):
        if prs.slides.part.related_part(sld_id.rId) if False else None:
            pass
    for sld_id in list(sld_ids):
        if prs.part.related_part(sld_id.get(qn("r:id"))) is slide.part:
            prs.part.drop_rel(sld_id.get(qn("r:id")))
            sld_ids.remove(sld_id)
            return


# ---------- data ----------
bench = json.loads((ROOT / "docs" / "benchmark.json").read_text(encoding="utf-8"))
total_runs = sum(b["runs"] for b in bench)
total_ok = sum(b["correct"] for b in bench)
lat_all = sum(b["latency_mean_s"] * b["runs"] for b in bench) / total_runs
tok_in = sum(b["tokens_in_mean"] * b["runs"] for b in bench) / total_runs
tok_out = sum(b["tokens_out_mean"] * b["runs"] for b in bench) / total_runs
PRICE_IN, PRICE_OUT = 0.30, 2.50          # USD per 1M tokens, Gemini 2.5 Flash list price
cost_claim = tok_in / 1e6 * PRICE_IN + tok_out / 1e6 * PRICE_OUT

prs = Presentation(str(HERE / "template.pptx"))
S = list(prs.slides)

# ---------- 1. Team details ----------
s = S[0]
tb = next(sh for sh in s.shapes if sh.has_text_frame and "Team Details" in sh.text_frame.text)
values = {"Team name: ": "ClaimGuard", "Team leader name: ": "Priyam Maini",
          "Problem Statement: ": "Banking, Financial Services and Insurance. Catching fraud and risk in insurance claims with AI"}
for p in tb.text_frame.paragraphs:
    if p.runs and p.runs[0].text in values:
        label = p.runs[0]
        r = copy.deepcopy(label._r)
        label._r.addnext(r)
        p.runs[1].text = values[label.text]
        p.runs[1].font.bold = False
        label.font.bold = True

# ---------- 3. Brief about the idea ----------
s = S[2]
set_title(s, "Brief about the idea")
headline(s, "Reads every photo and document, then finds where they disagree")
text(s, L, 1.85, 4.55, 2.6, [
    "Every insurance claim arrives as a story plus a pile of evidence: photos, invoices, estimates, "
    "police reports. Fraud rarely sits in one document. It hides in the gaps between them. A repair bill "
    "dated before the accident. Front-end parts billed for a rear-end crash. A phone reported lost two "
    "days before it was bought.",
    "ClaimGuard sends the whole claim to Gemini on Vertex AI in one request. It reads everything, checks each "
    "piece against the others and against the insurer's own policy rules, and gives the adjuster a risk "
    "score, a recommendation and the evidence behind every finding.",
    [("The adjuster still makes the call. ClaimGuard makes sure nothing gets missed.", {"bold": True, "color": INK})],
], size=11.5, space_after=7)
verdict = crop(SHOTS / "result2.png", (1040, 170, 2690, 668), "verdict2.png")
picture(s, verdict, 5.15, 1.9, w=4.5)
caption(s, 5.15, 3.32, 4.5, "Live result on a test claim: the photo shows front damage, the claimant says rear")
stats = [("$308.6B", "estimated yearly cost of insurance fraud in the US (Coalition Against Insurance Fraud, 2022)"),
         (f"{lat_all:.0f} s", "average time for a full multi-document review"),
         ("1 request", "all photos and documents reviewed together")]
for i, (big, small) in enumerate(stats):
    x = 5.15 + i * 1.53
    text(s, x, 3.85, 1.45, 0.45, big, size=18, color=INK, bold=True)
    text(s, x, 4.3, 1.42, 1.0, small, size=9, color=MUTED)

# ---------- 4. Opportunities ----------
s = S[3]
set_title(s, "Opportunities")
headline(s, "Most fraud tools look at numbers. ClaimGuard looks at the evidence")
cols = [
    ("How is it different from existing ideas?",
     ["Rule engines and scoring models work from structured fields like amount, date and claim history. "
      "They cannot open a PDF or look at a photo.",
      "Document AI tools pull fields out of each file, but check every file on its own.",
      "ClaimGuard compares every file with every other file and with the claimant's own words. That is "
      "where staged and padded claims fall apart."]),
    ("How will it solve the problem?",
     ["Adjusters get a first-pass review in about half a minute instead of reading every page.",
      "Honest claims with consistent evidence can be fast-tracked.",
      "Suspicious claims reach the investigations team with each contradiction already listed and "
      "traced back to its source file."]),
    ("USP of the proposed solution",
     [[("Explainable. ", {"bold": True, "color": INK}), "Every flag names the file and the exact detail behind it."],
      [("Hybrid scoring. ", {"bold": True, "color": INK}), "Hard policy rules set a minimum score the model cannot argue below."],
      [("Plug-in. ", {"bold": True, "color": INK}), "One REST API, installed inside the insurer's own Google Cloud project, so claim data never leaves it."]]),
]
cw, gap = (W - 2 * 0.25) / 3, 0.25
for i, (head, body) in enumerate(cols):
    x = L + i * (cw + gap)
    box(s, x, 1.8, cw, 3.5)
    text(s, x + 0.18, 1.98, cw - 0.36, 0.5, head, size=12.5, color=INK, bold=True)
    text(s, x + 0.18, 2.55, cw - 0.36, 2.7, body, size=10.5, space_after=6)

# ---------- 5. Features ----------
s = S[4]
set_title(s, "List of features offered by the solution")
headline(s, "Everything an adjuster needs for a first-pass review, in one screen")
feats = [
    ("01", "Reads the whole claim", "Photos, invoices, receipts, estimates and police reports, as images or PDFs, in a single Gemini request."),
    ("02", "Cross-checks the evidence", "Timelines, damage against repairs, invoice arithmetic, the statement against third-party reports."),
    ("03", "Policy rule floor", "Coverage limits, brand-new policies, late reporting and claim frequency. Thresholds are set per insurer."),
    ("04", "Explainable verdict", "A 0 to 100 risk score and approve, review or investigate, with every finding tied to its source file."),
    ("05", "Ask about the claim", "Follow-up questions and polite document requests to the claimant, answered only from that claim's record."),
    ("06", "Ready to integrate", "Authenticated /v1 REST API, OpenAPI docs, your own claim number carried through, printable reports."),
]
cw, ch, gx, gy = (W - 2 * 0.25) / 3, 1.6, 0.25, 0.22
for i, (num, head, body) in enumerate(feats):
    x = L + (i % 3) * (cw + gx)
    y = 1.8 + (i // 3) * (ch + gy)
    box(s, x, y, cw, ch)
    text(s, x + 0.18, y + 0.16, 0.5, 0.3, num, size=10, color=NAVY, bold=True)
    text(s, x + 0.18, y + 0.42, cw - 0.36, 0.3, head, size=12.5, color=INK, bold=True)
    text(s, x + 0.18, y + 0.77, cw - 0.36, 0.8, body, size=10)

# ---------- 6. Process flow ----------
s = S[5]
set_title(s, "Process flow diagram")
headline(s, "From a new claim to an adjuster decision in about half a minute")
y0 = 2.0
# step 1
boxed(s, L, y0 + 0.75, 1.75, 1.35, [[("Claim arrives", {"bold": True, "color": INK, "size": 11.5})],
      "Statement, policy details, photos and documents, from the app or the /v1 API"], size=9.5)
# parallel checks
boxed(s, 2.55, y0, 2.35, 1.3, [[("Gemini 2.5 Flash on Vertex AI", {"bold": True, "color": INK, "size": 11.5})],
      "Extracts facts from each file, describes the photos, compares every source with every other"],
      size=9.5, fill=RGBColor(0xE8, 0xEE, 0xF5))
boxed(s, 2.55, y0 + 1.6, 2.35, 1.3, [[("Policy rule engine", {"bold": True, "color": INK, "size": 11.5})],
      "Coverage limit, policy age, late reporting, claim frequency. Plain, testable code"], size=9.5)
# combine
boxed(s, 5.3, y0 + 0.75, 1.6, 1.35, [[("Final score", {"bold": True, "color": INK, "size": 11.5})],
      "The higher of the model score and the rule minimum"], size=9.5)
# outcomes
outs = [("Approve", "Fast-track a consistent claim", GREEN, RGBColor(0xEC, 0xF3, 0xEE)),
        ("Manual review", "An adjuster checks the open points", AMBER, RGBColor(0xF6, 0xF1, 0xE4)),
        ("Investigate", "Route to the fraud team with evidence", RED, RGBColor(0xF6, 0xEC, 0xEA))]
for i, (h, b, c, f) in enumerate(outs):
    yy = y0 - 0.05 + i * 1.02
    boxed(s, 7.35, yy, 2.3, 0.85, [[(h, {"bold": True, "color": c, "size": 11.5})], b], size=9.5, fill=f)
    arrow(s, 6.9, y0 + 1.42, 7.35, yy + 0.42)
arrow(s, L + 1.75, y0 + 1.42, 2.55, y0 + 0.65)
arrow(s, L + 1.75, y0 + 1.42, 2.55, y0 + 2.25)
arrow(s, 4.9, y0 + 0.65, 5.3, y0 + 1.3)
arrow(s, 4.9, y0 + 2.25, 5.3, y0 + 1.55)
text(s, L, 5.0, W, 0.3, "At every step the adjuster can open the evidence, ask ClaimGuard a follow-up question, "
     "or print the assessment for the claim file. ClaimGuard never approves or denies a claim on its own.",
     size=10, color=MUTED)

# ---------- 7. Mock / product surfaces ----------
s = S[6]
set_title(s, "Wireframes and mock diagrams of the proposed solution")
headline(s, "Three ways in: a product page, an adjuster app and an API")
landing = crop(SHOTS / "landing.png", (0, 0, 2880, 1800), "landing.png", max_w=1400)
app1 = crop(SHOTS / "result1.png", (0, 0, 2880, 1800), "result1.png", max_w=1400)
docs = crop(SHOTS / "docs.png", (0, 0, 2880, 1800), "docs.png", max_w=1400)
pw = (W - 2 * 0.2) / 3
for i, (img, cap) in enumerate([(landing, "Product page for insurers, at /"),
                                 (app1, "Adjuster app, at /app"),
                                 (docs, "Integration API reference, at /docs")]):
    x = L + i * (pw + 0.2)
    picture(s, img, x, 1.75, w=pw)
    caption(s, x, 1.75 + pw * 0.625 + 0.08, pw, cap)
text(s, L, 4.25, W, 1.0, [
    "The interface is deliberately plain. Adjusters see the score and recommendation first, then the findings "
    "with their sources, then the consistency checks, extracted fields and photo review. Nothing is hidden behind a click.",
], size=10.5)

# ---------- 8. Architecture ----------
s = S[7]
set_title(s, "Architecture diagram of the proposed solution")
headline(s, "One small service inside the insurer's own Google Cloud project")
# clients
boxed(s, L, 2.05, 1.8, 0.9, [[("Adjuster", {"bold": True, "color": INK, "size": 11})], "Browser, /app"], size=9.5)
boxed(s, L, 3.35, 1.8, 0.9, [[("Claims system", {"bold": True, "color": INK, "size": 11})], "Calls /v1 with an API key"], size=9.5)
# project boundary
box(s, 2.45, 1.75, 7.2, 3.15, fill=None, line=MUTED, radius=0.03)
text(s, 2.6, 1.82, 5, 0.3, "Insurer's Google Cloud project", size=9.5, color=MUTED, bold=True)
# cloud run
box(s, 2.7, 2.15, 3.55, 2.55, fill=RGBColor(0xE8, 0xEE, 0xF5), radius=0.04)
text(s, 2.85, 2.25, 3.3, 0.3, "Cloud Run service", size=11.5, color=INK, bold=True)
for i, (h, b) in enumerate([("Web app", "Product page and adjuster UI"),
                            ("Integration API", "/v1, API keys, OpenAPI docs"),
                            ("Rule engine", "Thresholds set per deployment")]):
    boxed(s, 2.85, 2.62 + i * 0.67, 3.25, 0.57, [[(h + "  ", {"bold": True, "color": INK}), b]],
          size=9.5, fill=WHITE, pad=0.1, anchor=MSO_ANCHOR.MIDDLE)
# vertex
boxed(s, 6.75, 2.15, 2.7, 1.25, [[("Vertex AI", {"bold": True, "color": INK, "size": 11.5})],
      "Gemini 2.5 Flash. Multimodal input, JSON schema output, called through the service account"],
      size=9.5, fill=RGBColor(0xE8, 0xEE, 0xF5))
boxed(s, 6.75, 3.6, 2.7, 1.1, [[("Cloud Build + Artifact Registry", {"bold": True, "color": INK, "size": 10.5})],
      "deploy.sh builds from source and rolls out a new revision"], size=9.5)
arrow(s, L + 1.8, 2.5, 2.7, 2.9)
arrow(s, L + 1.8, 3.8, 2.7, 3.55)
arrow(s, 6.25, 2.78, 6.75, 2.78)
arrow(s, 6.75, 4.15, 6.25, 4.15)
text(s, L, 5.0, W, 0.3, "No model keys in code: Cloud Run's service account holds roles/aiplatform.user. Claim files are "
     "processed in memory and never leave the project.", size=10, color=MUTED)

# ---------- 9. Technologies ----------
s = S[8]
set_title(s, "Technologies to be used in the solution")
headline(s, "Built on Google Cloud, with as few moving parts as possible")
rows = [("Layer", "Technology", "Why we chose it"),
        ("AI reasoning", "Gemini 2.5 Flash on Vertex AI", "Reads images and PDFs natively, follows a strict JSON schema, fast and cheap per claim"),
        ("Test data", "Gemini 2.5 Flash Image", "Realistic synthetic damage photos, so no real customer data is ever used"),
        ("Hosting", "Cloud Run", "Serverless, scales with claim volume, installs into any Google Cloud project"),
        ("Build", "Cloud Build, Artifact Registry", "Source to running container with one command, no local Docker"),
        ("Backend", "Python, FastAPI, google-genai SDK", "Typed endpoints with OpenAPI docs generated automatically"),
        ("Frontend", "HTML, CSS, plain JavaScript", "No build step, loads instantly, easy for an insurer's team to restyle"),
        ("Security", "IAM service account, API keys on /v1", "Model access through IAM, integrations through rotating keys")]
tbl = s.shapes.add_table(len(rows), 3, Inches(L), Inches(1.8), Inches(W), Inches(3.4)).table
tbl.columns[0].width, tbl.columns[1].width = Inches(1.4), Inches(2.9)
tbl.columns[2].width = Inches(W - 4.3)
for r, row in enumerate(rows):
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = INK if r == 0 else (CARD if r % 2 else WHITE)
        cell.margin_left = cell.margin_right = Inches(0.1)
        cell.margin_top = cell.margin_bottom = Inches(0.05)
        tf = cell.text_frame
        tf.text = ""
        run = tf.paragraphs[0].add_run()
        run.text = val
        run.font.name = FONT
        run.font.size = Pt(10 if r else 10)
        run.font.bold = r == 0 or c == 1
        run.font.color.rgb = WHITE if r == 0 else (INK if c < 2 else BODY)
tbl_el = tbl._tbl
tbl_el.tblPr.set("firstRow", "1")
tbl_el.tblPr.set("bandRow", "0")

# ---------- 10. Cost ----------
s = S[9]
set_title(s, "Estimated implementation cost")
headline(s, f"About {cost_claim * 100:.1f} cents of model cost per claim reviewed")
claims_month = 10000
model_month = cost_claim * claims_month
text(s, L, 1.85, 3.0, 0.7, f"${cost_claim:.4f}", size=34, color=INK, bold=True)
text(s, L, 2.55, 3.0, 0.9, f"Gemini 2.5 Flash cost per claim, measured over {total_runs} benchmark runs "
     f"(average {tok_in:,.0f} input and {tok_out:,.0f} output tokens)", size=10, color=MUTED)
text(s, L, 3.45, 3.0, 0.7, f"${model_month:,.0f}", size=34, color=INK, bold=True)
text(s, L, 4.15, 3.0, 0.9, f"model cost for {claims_month:,} claims a month, with no volume discounts", size=10, color=MUTED)
crow = [("Item", "Basis", "Monthly estimate"),
        ("Gemini 2.5 Flash on Vertex AI", f"$0.30 per 1M input, $2.50 per 1M output tokens; {claims_month:,} claims", f"${model_month:,.0f}"),
        ("Cloud Run", "1 vCPU, 1 GiB, one warm instance, autoscaling to 3", "$15 to $50"),
        ("Cloud Build, Artifact Registry", "A few deploys a month, one image", "Under $5"),
        ("Pilot total", "Shadow mode on a live claim feed", f"About ${model_month + 55:,.0f}")]
tbl = s.shapes.add_table(len(crow), 3, Inches(3.65), Inches(1.9), Inches(6.0), Inches(2.2)).table
for i, wd in enumerate((2.0, 2.75, 1.25)):
    tbl.columns[i].width = Inches(wd)
for r, row in enumerate(crow):
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.fill.solid()
        last = r == len(crow) - 1
        cell.fill.fore_color.rgb = INK if r == 0 else (RGBColor(0xE8, 0xEE, 0xF5) if last else (CARD if r % 2 else WHITE))
        cell.margin_left = cell.margin_right = Inches(0.08)
        tf = cell.text_frame
        tf.text = ""
        run = tf.paragraphs[0].add_run()
        run.text = val
        run.font.name = FONT
        run.font.size = Pt(9.5)
        run.font.bold = r == 0 or last or c == 0
        run.font.color.rgb = WHITE if r == 0 else INK
        if c == 2:
            tf.paragraphs[0].alignment = PP_ALIGN.RIGHT
tbl._tbl.tblPr.set("bandRow", "0")
text(s, 3.65, 4.35, 6.0, 0.9, [
    "Compare that with the cost of one missed fraudulent claim, often thousands of dollars. Catching a single "
    "padded repair bill like our $9,000 test case pays for months of running ClaimGuard.",
    [("Estimates use Google Cloud list prices at the time of writing and exclude taxes and committed-use discounts.",
      {"size": 8.5, "color": MUTED})]], size=10.5, space_after=6)

# ---------- 11. Snapshots ----------
s = S[10]
set_title(s, "Snapshots of the prototype")
headline(s, "A $9,000 claim flagged in 30 seconds, with the evidence laid out")
full = Image.open(SHOTS / "result2.png")
left = crop(SHOTS / "result2.png", (1040, 690, 2690, 1800), "findings2.png")
picture(s, verdict, L, 1.85, w=4.55)
picture(s, left, L + 4.75, 1.85, w=W - 4.75)
text(s, L, 3.35, 4.55, 1.9, [
    [("What it caught on this claim", {"bold": True, "color": INK})],
    "The photo shows front-end damage, but the claimant describes a hit to the back of the car.",
    "The invoice lists a hood, radiator and headlamps, and is dated three days before the incident.",
    "Its line items add up to $4,750, yet the total says $9,000.",
    "The same shop's own estimate for rear damage was $1,550.",
], size=10, space_after=5)

# ---------- 12. Performance ----------
s = S[11]
set_title(s, "Prototype performance report and benchmarking")
headline(s, f"{total_ok} of {total_runs} runs reached the expected decision")
callouts = [(f"{total_ok}/{total_runs}", "runs with the expected recommendation"),
            (f"{lat_all:.0f} s", "average review time, 2 to 3 files per claim"),
            (f"{max(b['score_max'] - b['score_min'] for b in bench)} pts", "largest score spread across repeat runs")]
for i, (big, small) in enumerate(callouts):
    x = L + i * 1.65
    text(s, x, 1.8, 1.6, 0.55, big, size=24, color=INK, bold=True)
    text(s, x, 2.35, 1.5, 0.6, small, size=9.5, color=MUTED)
prow = [("Test claim", "Files", "Expected", "Correct", "Score range", "Avg time")]
names = ["Consistent rear-end collision", "Inflated repair invoice", "Phone lost before purchase", "Late water damage report", "Photo of a different car"]
for n, b in zip(names, bench):
    prow.append((n, str(b["files"]), b["expected"].replace("_", " ").title(), f"{b['correct']}/{b['runs']}",
                 f"{b['score_min']} to {b['score_max']}", f"{b['latency_mean_s']} s"))
tbl = s.shapes.add_table(len(prow), 6, Inches(L), Inches(2.95), Inches(W), Inches(1.85)).table
for i, wd in enumerate((2.9, 0.7, 1.6, 1.0, 1.6, 1.5)):
    tbl.columns[i].width = Inches(wd)
for r, row in enumerate(prow):
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = INK if r == 0 else (CARD if r % 2 else WHITE)
        cell.margin_left = cell.margin_right = Inches(0.08)
        tf = cell.text_frame
        tf.text = ""
        run = tf.paragraphs[0].add_run()
        run.text = val
        run.font.name = FONT
        run.font.size = Pt(10)
        run.font.bold = r == 0 or c == 0
        run.font.color.rgb = WHITE if r == 0 else INK
tbl._tbl.tblPr.set("bandRow", "0")
text(s, 5.35, 1.8, 4.3, 1.2, [
    f"Each claim was run {bench[0]['runs']} times through the production code path (Gemini 2.5 Flash on Vertex AI, temperature 0.2), "
    "four at a time. Scores move a few points between runs, but the recommendation holds, and the rule floor "
    "keeps hard policy breaches from ever dropping into low risk."], size=10)
text(s, L, 5.08, W, 0.3, "Test set: five synthetic claims with known answers. Next step is a labelled set of real, "
     "anonymised claims from a pilot partner.", size=9.5, color=MUTED)

# ---------- 13. Future ----------
s = S[12]
set_title(s, "Additional details and future development")
headline(s, "From a working prototype to a pilot with a real insurer")
phases = [("Next 30 days", "Pilot ready",
           ["Firestore for claim history and an audit log", "Pub/Sub intake from claim portals", "Single sign-on for adjusters"]),
          ("Next quarter", "Network signals",
           ["BigQuery history to catch reused invoices and repeat vendors", "Fraud ring detection across claims",
            "Vertex AI model tuned on real investigation outcomes"]),
          ("Later", "More lines of business",
           ["Photo forensics: metadata checks, reused or stock images", "Health and travel claims",
            "KYC and loan document checks for banks"])]
cw = (W - 2 * 0.25) / 3
for i, (when, what, items) in enumerate(phases):
    x = L + i * (cw + 0.25)
    box(s, x, 1.8, cw, 2.75)
    text(s, x + 0.18, 1.95, cw - 0.36, 0.3, when.upper(), size=9, color=NAVY, bold=True)
    text(s, x + 0.18, 2.22, cw - 0.36, 0.35, what, size=13, color=INK, bold=True)
    text(s, x + 0.18, 2.7, cw - 0.36, 1.8, items, size=10, space_after=7)
text(s, L, 4.75, W, 0.6, [[("Business model. ", {"bold": True, "color": INK}),
     "Priced per claim reviewed and installed in the customer's own cloud account, so insurers keep control of their "
     "data and pay in line with their claim volume."]], size=10.5)

# ---------- 14. Links ----------
s = S[13]
set_title(s, "Provide links to your:")
links = [("GitHub public repository", REPO), ("Demo video (3 minutes)", VIDEO), ("Final product link", LIVE),
         ("API reference", LIVE + "/docs")]
for i, (lab, url) in enumerate(links):
    y = 1.4 + i * 0.85
    box(s, L, y, W, 0.68)
    text(s, L + 0.2, y + 0.1, 3.0, 0.5, lab, size=12, color=INK, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    tb = text(s, L + 3.1, y + 0.1, W - 3.3, 0.5, url, size=12, color=NAVY if url.startswith("http") else RED,
              anchor=MSO_ANCHOR.MIDDLE)
    if url.startswith("http"):
        tb.text_frame.paragraphs[0].runs[0].hyperlink.address = url

# ---------- remove instruction slide and spare blank slide ----------
delete_slide(prs, S[14])
delete_slide(prs, S[1])

out = HERE / "ClaimGuard_AI_Builder_Cup.pptx"
prs.save(out)
print("saved", out, f"| cost/claim ${cost_claim:.4f} | {total_ok}/{total_runs} correct | {lat_all:.1f}s")
